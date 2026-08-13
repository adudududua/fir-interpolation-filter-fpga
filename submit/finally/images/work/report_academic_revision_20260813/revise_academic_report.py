from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def find_exact(doc: Document, text: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}; found {len(matches)}")
    return matches[0]


def find_contains(doc: Document, fragment: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if fragment in p.text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph containing {fragment!r}; found {len(matches)}")
    return matches[0]


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    style = paragraph.style
    alignment = paragraph.alignment
    fmt = paragraph.paragraph_format
    saved = {
        "left_indent": fmt.left_indent,
        "right_indent": fmt.right_indent,
        "first_line_indent": fmt.first_line_indent,
        "space_before": fmt.space_before,
        "space_after": fmt.space_after,
        "line_spacing": fmt.line_spacing,
        "keep_with_next": fmt.keep_with_next,
        "keep_together": fmt.keep_together,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
    }
    first = paragraph.runs[0] if paragraph.runs else None
    font = None
    if first is not None:
        font = {
            "name": first.font.name,
            "size": first.font.size,
            "bold": first.font.bold,
            "italic": first.font.italic,
            "underline": first.font.underline,
        }
    paragraph.clear()
    run = paragraph.add_run(text)
    if font:
        for key, value in font.items():
            if value is not None:
                setattr(run.font, key, value)
        if font.get("name"):
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font["name"])
    paragraph.style = style
    paragraph.alignment = alignment
    for key, value in saved.items():
        if value is not None:
            setattr(paragraph.paragraph_format, key, value)


def paragraph_before(anchor: Paragraph, text: str, style: str) -> Paragraph:
    paragraph = anchor.insert_paragraph_before(text, style=style)
    return paragraph


def move_before(anchor: Paragraph, element) -> None:
    anchor._p.addprevious(element)


def drawing_before(caption: Paragraph):
    element = caption._p.getprevious()
    while element is not None:
        if element.tag == qn("w:p") and element.xpath(".//w:drawing"):
            return element
        element = element.getprevious()
    raise RuntimeError(f"Drawing before caption not found: {caption.text}")


def remove_block(start_element, stop_element) -> None:
    parent = start_element.getparent()
    element = start_element
    while element is not None and element is not stop_element:
        following = element.getnext()
        parent.remove(element)
        element = following
    if element is not stop_element:
        raise RuntimeError("Stop element not reached while removing chapter block")


def set_cell_text(cell, text: str, size: float = 7.2, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def create_table_before(anchor: Paragraph, doc: Document, rows: list[list[str]], widths: list[float]) -> None:
    if any(len(row) != len(rows[0]) for row in rows):
        raise RuntimeError("Ragged table data")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            set_cell_text(table.cell(ri, ci), value, 6.9 if ri else 7.1, ri == 0)
            table.cell(ri, ci).width = Inches(widths[ci])
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    table._tbl.getparent().remove(table._tbl)
    move_before(anchor, table._tbl)


def delete_table_rows(table, keep_indices: list[int]) -> None:
    keep = set(keep_indices)
    for index in reversed(range(len(table.rows))):
        if index not in keep:
            table._tbl.remove(table.rows[index]._tr)


def set_heading_text_preserve_bookmarks(paragraph: Paragraph, text: str) -> None:
    if not paragraph.style.name.startswith("Heading"):
        raise RuntimeError(f"Not a heading: {paragraph.text}")
    first_text = paragraph._p.find(qn("w:r"))
    for run in list(paragraph._p.findall(qn("w:r"))):
        paragraph._p.remove(run)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    bookmark_end = paragraph._p.find(qn("w:bookmarkEnd"))
    if bookmark_end is not None:
        paragraph._p.insert(paragraph._p.index(bookmark_end), run)
    else:
        paragraph._p.append(run)


def add_toc_bookmark(paragraph: Paragraph, bookmark_id: int, name: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    first_run = paragraph._p.find(qn("w:r"))
    if first_run is None:
        raise RuntimeError("Heading has no run")
    paragraph._p.insert(paragraph._p.index(first_run), start)
    paragraph._p.append(end)


def make_toc_entry(template: Paragraph, title: str, page: str, bookmark: str):
    new_p = OxmlElement("w:p")
    ppr = template._p.find(qn("w:pPr"))
    if ppr is not None:
        from copy import deepcopy

        new_p.append(deepcopy(ppr))
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    hyperlink.set(qn("w:history"), "1")

    text_run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = title
    text_run.append(text_node)
    tab = OxmlElement("w:tab")
    text_run.append(tab)
    hyperlink.append(text_run)

    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    hyperlink.append(begin_run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark} \\h "
    instr_run.append(instr)
    hyperlink.append(instr_run)

    sep_run = OxmlElement("w:r")
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(sep)
    hyperlink.append(sep_run)

    page_run = OxmlElement("w:r")
    page_text = OxmlElement("w:t")
    page_text.text = page
    page_run.append(page_text)
    hyperlink.append(page_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    hyperlink.append(end_run)
    new_p.append(hyperlink)
    return new_p


def patch_toc_entry(source: Paragraph, title: str, page: str, bookmark: str) -> Paragraph:
    new_p = make_toc_entry(source, title, page, bookmark)
    source._p.addnext(new_p)
    return Paragraph(new_p, source._parent)


def rewrite_toc_visible_title(paragraph: Paragraph, title: str) -> None:
    hyperlink = paragraph._p.find(qn("w:hyperlink"))
    if hyperlink is None:
        raise RuntimeError(f"TOC hyperlink not found: {paragraph.text}")
    first_text = None
    before_field = True
    for child in hyperlink:
        field = child.find(qn("w:fldChar"))
        if field is not None and field.get(qn("w:fldCharType")) == "begin":
            before_field = False
            continue
        if not before_field:
            continue
        for text_node in child.iter(qn("w:t")):
            if first_text is None:
                first_text = text_node
                text_node.text = title
            else:
                text_node.text = ""
    if first_text is None:
        raise RuntimeError(f"TOC visible title node not found: {paragraph.text}")


def set_update_fields(path: Path) -> None:
    temporary = path.with_suffix(".updatefields.tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/settings.xml":
                if b"<w:updateFields" in data:
                    data = re.sub(rb'<w:updateFields[^>]*/>', b'<w:updateFields w:val="true"/>', data, count=1)
                else:
                    data = data.replace(b"</w:settings>", b'<w:updateFields w:val="true"/></w:settings>')
            target.writestr(item, data)
    temporary.replace(path)


def normalize_references(doc: Document) -> None:
    references = [
        "[1] AMD. FIR Compiler v7.2 LogiCORE IP product guide (PG149)[EB/OL]. (2025-12-17)[2026-07-08]. https://docs.amd.com/r/en-US/pg149-fir-compiler/Support-Resources.",
        "[2] CANESE L, CARDARILLI G C, DI NUNZIO L, et al. Efficient digital implementation of a multirate-based variable fractional delay filter for wideband beamforming[J]. IEEE Transactions on Circuits and Systems II: Express Briefs, 2023, 70(6): 2231-2235. DOI: 10.1109/TCSII.2023.3236066.",
        "[3] DATTA D, DUTTA H S. Implementation of polyphase digital down converter for wireless applications[J]. Microprocessors and Microsystems, 2023, 100: 104850. DOI: 10.1016/j.micpro.2023.104850.",
        "[4] DATTA D, DUTTA H S. Hardware optimized digital down converter for multi-standard radio receiver[J]. Analog Integrated Circuits and Signal Processing, 2024, 118(3): 567-575. DOI: 10.1007/s10470-023-02227-y.",
        "[5] GADAWE N T, HAMAD R W, QADDOORI S L. Efficient implementation of fractional sampling rate conversion (SRC) on FPGA[J]. Diagnostyka, 2025, 26(3): 2025308. DOI: 10.29354/diag/208854.",
        "[6] GALINDO GUARCH F J, BAUDRENGHIEN P, MORENO AROSTEGUI J M. An architecture for real-time arbitrary and variable sampling rate conversion with application to the processing of harmonic signals[J]. IEEE Transactions on Circuits and Systems I: Regular Papers, 2020, 67(5): 1653-1666. DOI: 10.1109/TCSI.2019.2960686.",
        "[7] MOTTA L L, ACUÑA ACURIO B A, ANICETO N F T, et al. Design and implementation of a digital down/up conversion directly from/to RF channels in HDL[J]. Integration, 2019, 68: 30-37. DOI: 10.1016/j.vlsi.2019.05.006.",
        "[8] PINJERLA S, RAO S S, REDDY P C. Design of energy efficient and reconfigurable sample rate converter using FPGA devices[J]. Indonesian Journal of Electrical Engineering and Computer Science, 2024, 36(2): 854-862. DOI: 10.11591/ijeecs.v36.i2.pp854-862.",
        "[9] PINJERLA S, RAO S S, REDDY P C. Multi-stage decimation with hybrid CIC-polyphase filtering for IoT gateway sample rate conversion[J]. Scientific Reports, 2026, 16(1): 2016. DOI: 10.1038/s41598-025-31617-7.",
        "[10] SHAHABUDDIN S, MANNINEN P, JUNTTI M. A fractional sample rate converter with parallelized multiphase output: algorithm and FPGA implementation[J]. Journal of Signal Processing Systems, 2022, 94(12): 1459-1469. DOI: 10.1007/s11265-022-01776-1.",
        "[11] ZEINEDDINE A, NAFKHA A, PAQUELET S, et al. Comprehensive survey of FIR-based sample rate conversion[J]. Journal of Signal Processing Systems, 2021, 93(1): 113-125. DOI: 10.1007/s11265-020-01575-6.",
        "[12] HOGENAUER E B. An economical class of digital filters for decimation and interpolation[J]. IEEE Transactions on Acoustics, Speech, and Signal Processing, 1981, 29(2): 155-162.",
        "[13] CROCHIERE R E, RABINER L R. Multirate digital signal processing[M]. Englewood Cliffs: Prentice-Hall, 1983.",
        "[14] HARRIS F J. Multirate signal processing for communication systems[M]. Upper Saddle River: Prentice Hall PTR, 2004.",
        "[15] XILINX. 7 series DSP48E1 slice user guide (UG479)[R]. 2018.",
        "[16] XILINX. 7 series FPGAs configurable logic block user guide (UG474)[R]. 2016.",
        "[17] AMD. Vivado Design Suite user guide: design analysis and closure techniques (UG906), Vivado 2025.2[R]. 2025.",
        "[18] 全国大学生集成电路创新创业大赛. 高阶数字插值滤波器设计与验证赛题及全国总决赛技术要求[Z]. 2026.",
    ]
    heading = find_exact(doc, "参考文献")
    paragraphs = list(doc.paragraphs)
    start = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is heading._p) + 1
    for offset, text in enumerate(references):
        paragraph = paragraphs[start + offset]
        if not paragraph.text.strip().startswith(f"[{offset + 1}]"):
            raise RuntimeError(f"Reference index mismatch at {offset + 1}: {paragraph.text}")
        set_paragraph_text(paragraph, text)


def academic_normalization(doc: Document) -> dict[str, int]:
    replacements = {
        "Smoke/Release 17/17": "快速RTL回归与完整RTL回归均为17/17项满足判据",
        "Smoke和Release两种17项RTL回归": "快速RTL回归与完整RTL回归两组各17项测试",
        "RTL Smoke 17/17": "快速RTL回归17/17项满足判据",
        "board-verified版本": "经物理板功能验证的版本",
        "board-verified": "经物理板功能验证",
        "板测通过": "物理板功能验证结果满足判据",
        "板测版本": "物理板功能验证版本",
        "实板通过": "物理板功能验证结果满足判据",
        "实板验证": "物理板功能验证",
        "黄金模型": "位精确参考模型",
        "MATLAB golden": "MATLAB位精确参考模型",
        "对拍0 LSB": "逐样本比较的最大绝对误差为0 LSB",
        "生产配置": "目标实现配置",
    }
    counts = {key: 0 for key in replacements}
    appendix = find_exact(doc, "附录A 最终参数与数据格式速查")
    paragraphs = list(doc.paragraphs)
    appendix_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is appendix._p)
    for paragraph in paragraphs[:appendix_index]:
        if paragraph.style.name.lower().startswith("toc"):
            continue
        old = paragraph.text
        new = old
        for source, target in replacements.items():
            if source in new:
                counts[source] += new.count(source)
                new = new.replace(source, target)
        if new != old:
            if paragraph.style.name.startswith("Heading"):
                set_heading_text_preserve_bookmarks(paragraph, new)
            else:
                set_paragraph_text(paragraph, new)
    body = appendix._p.getparent()
    appendix_body_index = body.index(appendix._p)
    for table in doc.tables:
        if table._tbl.getparent() is not body or body.index(table._tbl) >= appendix_body_index:
            continue
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    old = paragraph.text
                    new = old
                    for source, target in replacements.items():
                        if source in new:
                            counts[source] += new.count(source)
                            new = new.replace(source, target)
                    if new != old:
                        set_paragraph_text(paragraph, new)
    return {key: value for key, value in counts.items() if value}


def targeted_academic_rewrites(doc: Document) -> None:
    replacements = {
        "Vivado 2025.2完整板级post-route占用218 LUT": "Vivado 2025.2完整板级布局布线后（post-route）占用218 LUT",
        "Medium置信度的vectorless估计": "Medium置信度的无活动向量（vectorless）估计",
        "通过路由后六模式仿真与用户物理板复测，证明双采样率族、各倍率档位和DAC输出链路均稳定工作": "通过布局布线后六模式仿真与物理板复测，验证双采样率族、各倍率档位和DAC输出链路的功能稳定性",
        "作为发布结论": "作为报告结论",
        "按照正式发布配置": "按照正式归档配置",
        "该波形仅用于证明调度与valid关系": "该波形仅用于验证调度与数据有效指示（valid）关系",
        "218-LUT正式配置采用24/20/20自有位精确参考模型": "218-LUT正式配置采用24/20/20位精确参考模型",
        "本轮RTL Smoke重新执行17项独立测试并全部通过。历史Release 17/17和14组全链覆盖继续作为既有发布证据，但不与本轮Smoke混写为同一次执行。": "本轮重新执行快速RTL回归（Smoke配置）的17项独立测试，全部满足判据。既有完整RTL回归（Release配置）17/17项和14组全链覆盖继续作为归档证据，但不与本轮快速回归表述为同一次执行。",
        "不可变板测标签快照": "经物理板功能验证的不可变标签快照",
        "与该板测标签的mode2参数精确匹配": "与该物理板验证标签的mode2参数精确匹配",
        "vectorless总功耗": "无活动向量（vectorless）总功耗",
        "当前发布点为218 LUT/365 FF/4 DSP/2 BRAM Tile": "LUT优先竞赛端点为218 LUT/365 FF/4 DSP/2 BRAM Tile",
        "239-LUT/3-DSP板测端点": "239-LUT/3-DSP物理板验证端点",
        "0.271 W仅作为vectorless估计": "0.271 W仅作为无活动向量（vectorless）估计",
    }
    for fragment, replacement in replacements.items():
        matches = [p for p in doc.paragraphs if fragment in p.text and not p.style.name.lower().startswith("toc")]
        if len(matches) != 1:
            raise RuntimeError(f"Expected one academic rewrite match for {fragment!r}; found {len(matches)}")
        paragraph = matches[0]
        set_paragraph_text(paragraph, paragraph.text.replace(fragment, replacement))
    for paragraph in doc.paragraphs:
        if paragraph.style.name.lower().startswith("toc") or "PASS关键字" not in paragraph.text:
            continue
        set_paragraph_text(paragraph, paragraph.text.replace("PASS关键字", "日志终止标志“PASS”"))


def build_revised_document(source: Path, output: Path, backup: Path) -> dict:
    backup.parent.mkdir(parents=True, exist_ok=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        shutil.copy2(source, backup)
    doc = Document(source)

    # Capture selected chapter assets before restructuring.
    old_tables = list(doc.tables)
    t34, t36, t37, t38 = (old_tables[index] for index in (34, 36, 37, 38))
    t44, t48, t49, t54, t58, t59, t60 = (old_tables[index] for index in (44, 48, 49, 54, 58, 59, 60))
    figure_assets = {}
    for caption_text in (
        "图9-1 两套128倍插值架构及最终实板端点",
        "图9-4 Vivado 2025.2同工具链LUT优化演进",
        "图9-6 固定4 RAMB18E1条件下的DSP-LUT Pareto与2-DSP候选复验",
        "图9-7 218-LUT与239-LUT板测版本的同口径资源比较",
    ):
        caption = find_exact(doc, caption_text)
        figure_assets[caption_text] = drawing_before(caption)

    # Strengthen related work without renumbering existing Chapter 1 sections.
    chapter2 = find_exact(doc, "第2章 插值理论与体系结构选择")
    review_heading = paragraph_before(chapter2, "1.4 国内外研究进展与本文技术定位", "Heading 2")
    max_bookmark_id = max(int(node.get(qn("w:id"))) for node in doc.element.body.iter(qn("w:bookmarkStart")) if (node.get(qn("w:id")) or "").isdigit())
    review_bookmark = "_TocAcademicRevision10104"
    add_toc_bookmark(review_heading, max_bookmark_id + 1, review_bookmark)
    review_paragraphs = [
        "多速率信号处理已形成系统的理论框架。Crochiere与Rabiner、Harris分别从抽取与插值、滤波器组及通信系统实现等角度论述了采样率转换原理[13-14]；Hogenauer提出的CIC滤波器以积分器—梳状器级联实现无乘法器的整数倍率变换，为高倍率插值的低资源实现奠定基础[12]。Zeineddine等对FIR型采样率转换方法进行了综述[11]。这些研究表明，多级分解、多相表示与CIC结构是降低高倍率采样率转换运算量的主要理论途径。",
        "近年来，研究重点由固定整数倍率扩展到分数、任意与可变采样率转换。已有工作分别讨论了宽带波束形成中的多速率可变分数时延[2]、FPGA分数采样率转换[5]、面向谐波信号的任意可变采样率架构[6]以及并行多相输出的分数采样率转换器[10]。在数字射频前端领域，多相数字下变频、面向多标准接收机的硬件优化以及HDL直接数字上/下变频也得到研究[3-4,7]。相关结果支持以多相分解和并行或时分复用结构在吞吐率、可配置性与硬件复杂度之间进行折衷。",
        "FPGA实现研究进一步关注低功耗可重构采样率转换[8]以及CIC—多相混合多级结构[9]。厂商FIR Compiler提供参数化FIR实现基线[1]；DSP48E1、可配置逻辑块和Vivado实现分析文档分别给出算术映射、逻辑资源与实现收敛的器件和工具依据[15-17]。由于器件型号、时钟约束、工具版本、顶层边界和定点格式都会影响实现结果，不同文献中的资源数据不宜直接进行数值比较，而应在统一实验条件下开展受控对照。",
        "就本文所列文献范围而言，既有研究已覆盖多相FIR、CIC、分数或可变采样率转换及多类FPGA硬件实现，但较少同时讨论双音频采样率族、128倍整数插值、严格线性相位、中间倍率可观测、有限字长误差、专用资源约束与物理板功能验证。依据赛题要求[18]，本文选择三级2倍多相FIR与16倍CIC级联，并将CIC通带补偿折叠到第三级FIR；在统一工具版本与统计边界下比较厂商IP基线、手写RTL及不同DSP映射，再以数值模型、位精确RTL、布局布线后报告和物理板功能验证构成分层证据。本文的研究重点不是提出新的采样率转换理论，而是在给定器件与指标约束下完成结构、字长、存储和调度的协同实现及可复核评价。",
    ]
    for text in review_paragraphs:
        paragraph_before(chapter2, text, "Normal")

    # Add the new entry to the cached TOC. Word/WPS will refresh page numbers on open.
    toc_13 = find_exact(doc, "1.3 最终结果与证据链\t9")
    patch_toc_entry(toc_13, "1.4 国内外研究进展与本文技术定位", "9", review_bookmark)

    chapter8 = find_exact(doc, "第8章 历史厂商IP基线与当前手写RTL的口径化对比")
    chapter10 = find_exact(doc, "第10章 创新点与工程价值")

    # Chapter 8: four evidence tables, no repeated figures.
    paragraph_before(chapter8, "第8章 历史厂商IP基线与手写RTL的口径化比较", "Heading 1")
    paragraph_before(chapter8, "本章以历史厂商IP为资源映射基线，并将其与手写RTL进行分组比较。由于工具版本、顶层功能和统计范围并非全部一致，分析仅在边界相同的实验组内计算差值；跨版本结果用于说明技术路线，不用于宣称直接资源收益。", "Normal")
    paragraph_before(chapter8, "8.1 比较目的与对象", "Heading 2")
    paragraph_before(chapter8, "比较对象分为三组：A组限定为155抽头4倍FIR前级，B组覆盖七级2倍FIR完整链，C组比较历史FIR—CIC IP路线与Vivado 2025.2下的手写FIR—CIC实现。A、B组用于评价通用IP映射与专用结构复用的差异；C组仅作为架构演进参照。", "Normal")
    paragraph_before(chapter8, "8.2 统一比较口径与证据边界", "Heading 2")
    paragraph_before(chapter8, "为避免将工具迁移、顶层封装和RTL优化的影响混合归因，本文仅在器件、功能边界、时钟约束、工具版本及统计口径一致时计算资源相对变化。A、B组采用Vivado 2018.3同口径结果；历史I-C与当前R-C因工具版本和完整系统边界不同，仅用于说明实现路线，不计算直接百分比。当前设计的定量优化收益由Vivado 2025.2同一完整顶层下的连续实现节点给出。", "Normal")
    paragraph_before(chapter8, "表8-1 三组IP与手写RTL的比较边界", "Caption")
    move_before(chapter8, t34._tbl)
    set_cell_text(t34.cell(3, 2), "R-C：218-LUT/4-DSP LUT优先端点（2025.2）")
    paragraph_before(chapter8, "8.3 历史IP与手写RTL的同口径结果", "Heading 2")
    paragraph_before(chapter8, "8.3.1 155抽头4倍前级", "Heading 3")
    paragraph_before(chapter8, "A组在相同的155抽头4倍前级功能边界内比较FIR Compiler与手写对称复用结构。两者均接收24 bit输入，但内部流水与乘加组织不同；因此本组结论仅对应前级模块，不外推到后续插值链或板级控制。", "Normal")
    paragraph_before(chapter8, "表8-2 155抽头4倍前级的同口径资源比较", "Caption")
    move_before(chapter8, t36._tbl)
    paragraph_before(chapter8, "手写前级相较厂商IP减少1268个LUT、2262个FF和41个DSP48E1，BRAM Tile不变。该结果表明，在固定系数与确定采样节拍条件下，对称性利用和串行乘加复用能够降低长FIR前级的专用乘法器与寄存器开销；其适用范围限于本组给定结构和实现约束。", "Normal")
    paragraph_before(chapter8, "8.3.2 七级2倍FIR完整链", "Heading 3")
    paragraph_before(chapter8, "B组比较相同七级2倍FIR任务下的厂商IP与逐级专用手写RTL。手写方案利用严格半带、逐级抽头裁剪、共享DSP与历史存储映射重新分配通用逻辑和专用资源。", "Normal")
    paragraph_before(chapter8, "表8-3 七级2倍FIR完整链的同口径资源比较", "Caption")
    move_before(chapter8, t37._tbl)
    paragraph_before(chapter8, "手写链以增加285个LUT为代价，减少1563个FF、6个DSP48E1和10.5个BRAM Tile，且在该实现条件下保持正时序裕量。该结果体现的是LUT与FF、DSP、BRAM之间的资源重分配，不宜简化为单一资源维度上的绝对优劣。功耗数据来自无活动向量估计，仅用于同流程趋势比较。", "Normal")
    paragraph_before(chapter8, "8.3.3 FIR—CIC路线的跨版本参照", "Heading 3")
    paragraph_before(chapter8, "C组用于观察FIR—CIC路线由历史IP组合向手写专用链的演进。由于两者采用不同Vivado版本、顶层封装与统计范围，表中只列绝对结果和口径说明，不计算百分比。", "Normal")
    paragraph_before(chapter8, "表8-4 FIR—CIC路线的跨版本实现参照", "Caption")
    move_before(chapter8, t38._tbl)
    set_cell_text(t38.cell(0, 2), "R-C：Vivado 2025.2 LUT优先端点")
    set_cell_text(t38.cell(3, 2), "4")
    paragraph_before(chapter8, "历史对照说明，三级2倍FIR与16倍CIC为后续专用化提供了可行架构，但当前资源收益只能依据Vivado 2025.2同一完整顶层下的292→221→218 LUT连续节点归因。218 LUT、239 LUT与268 LUT分别构成LUT优先、资源折中和DSP受限三类端点。", "Normal")
    paragraph_before(chapter8, "8.4 结果讨论与适用边界", "Heading 2")
    paragraph_before(chapter8, "A组的主要收益来自固定系数对称性与乘加复用；B组进一步表明，逐级专用设计可用有限LUT增量换取显著的FF、DSP和BRAM下降；C组则说明FIR—CIC混合路线具有继续压缩高采样率尾部状态的结构潜力。三组结论共同支持“按采样率与资源类型进行分层映射”，但不构成手写RTL在任意任务上均优于厂商IP的普遍性结论。", "Normal")
    paragraph_before(chapter8, "时序结论仅采用对应实现的布局布线后报告；功耗结论中的0.271 W为无活动向量估计且置信度为Medium，不能外推为实测板级功耗。厂商IP在参数配置、接口规范与验证效率方面仍具有工程价值，手写RTL则适合在系数、节拍和器件资源固定时开展针对性复用。", "Normal")
    paragraph_before(chapter8, "8.5 本章小结", "Heading 2")
    paragraph_before(chapter8, "历史对照表明，固定系数与确定采样节拍为对称化、时分复用和专用存储映射提供了实现空间。然而，该结论并不等同于手写RTL在任意条件下均优于厂商IP；其有效范围限于本文给定器件、约束和统计边界。当前实现的资源与时序结论以Vivado 2025.2同口径实现结果为准。", "Normal")

    # Chapter 9: retain only the evidence needed for architecture selection.
    paragraph_before(chapter8, "第9章 架构演进、资源权衡与实现选择", "Heading 1")
    paragraph_before(chapter8, "9.1 评价准则与证据层级", "Heading 2")
    paragraph_before(chapter8, "本文采用两级决策准则。首先，以双采样率128倍输出、通带、阻带、线性相位、中间节点、时序和板级功能构成可行性约束；随后在可行方案中比较LUT、FF、DSP、BRAM、验证成熟度及回退成本。由此得到的是给定器件与已覆盖候选集合下的Pareto选择，而非脱离约束的全局最优解。", "Normal")
    paragraph_before(chapter8, "表9-1 三种候选架构的统一设计与评价条件", "Caption")
    move_before(chapter8, t44._tbl)
    set_cell_text(t44.cell(6, 2), "完整板级布局布线后结果优先")
    paragraph_before(chapter8, "数学性能按双采样率六工况统一评价；硬件证据分为2018.3历史架构对照、Vivado 2025.2结构优化、有限字长Pareto和matched-OOC核心审计四层。不同层级用于回答不同问题，不能跨统计边界合并为单一百分比。", "Normal")
    paragraph_before(chapter8, "9.2 候选架构的横向比较", "Heading 2")
    paragraph_before(chapter8, "9.2.1 结构与性能", "Heading 3")
    create_table_before(
        chapter8,
        doc,
        [
            ["维度", "方案A：4倍FIR＋五级2倍", "方案B：七级2倍FIR", "方案C：三级2倍FIR＋CIC16"],
            ["倍率与滤波器", "4×2×2×2×2×2；155抽头前级", "2×七级；105/17/11/7/7/7/7抽头", "2×2×2×16；第三级折叠CIC补偿"],
            ["主要优势", "双采样率兼容；群时延较低；结构直观", "阻带余量高；各级规则；易于逐级验证", "双采样率；通用逻辑占用低；资源映射可调"],
            ["主要约束", "长前级资源集中；阻带余量较小", "高采样率尾部状态和量化边界累积", "调度、有限字长、CIC模运算和切换控制耦合"],
            ["设计取向", "通用兼容与时延", "频谱裕量与规则性", "资源效率与可复核实现"],
        ],
        [1.0, 1.55, 1.55, 1.65],
    )
    paragraph_before(chapter8, "表9-2 三种架构的结构特征与设计取向", "Caption")
    paragraph_before(chapter8, "方案A结构直观且群时延较低，但155抽头前级造成资源集中，阻带余量为0.33918 dB；方案B通过逐级专用设计取得78.61966 dB阻带衰减，频谱裕量最高，但高采样率尾部仍累积状态与量化边界；方案C采用三级2倍FIR与16倍CIC，在满足六工况频响及线性相位要求的前提下降低通用逻辑占用，同时提高了调度、有限字长与验证的耦合程度。三者分别对应通用兼容、频谱裕量和资源效率三类设计取向。", "Normal")
    paragraph_before(chapter8, "表9-3 三种架构的统一数学性能", "Caption")
    move_before(chapter8, t48._tbl)
    set_cell_text(t48.cell(0, 3), "方案C（本文实现）")
    paragraph_before(chapter8, "三种方案均满足基本频响与线性相位判据。方案B具有最大的阻带余量；方案C的六工况最差绝对通带偏差为0.007844 dB、最差阻带衰减为72.355 dB，并通过整数冲激镜像误差0 LSB验证严格线性相位。选型需在性能余量与实现资源之间联合判断。", "Normal")
    paragraph_before(chapter8, "9.2.2 资源口径与适用条件", "Heading 3")
    paragraph_before(chapter8, "表9-4 三种架构的分层资源证据", "Caption")
    move_before(chapter8, t49._tbl)
    paragraph_before(chapter8, "表9-4中的历史独立链、阶段性板级结果与当前完整系统不属于同一边界。它们用于说明架构瓶颈如何迁移；正式资源取舍以Vivado 2025.2连续节点和同流程OOC结果为准。", "Normal")
    paragraph_before(chapter8, "9.3 架构演进与实现优化", "Heading 2")
    paragraph_before(chapter8, "9.3.1 从全FIR到FIR—CIC", "Heading 3")
    paragraph_before(chapter8, "架构演进遵循“瓶颈定位—单变量候选—位精确验证—物理实现—板级回归”的受控流程。方案A的资源主要集中于155抽头前级及通用后级，因而演进为逐级专用的七级2倍FIR；方案B的高采样率尾部仍重复保存状态并形成多个量化边界，因而进一步以16倍CIC替代后四级FIR。各阶段仅在各自一致口径内解释。", "Normal")
    move_before(chapter8, figure_assets["图9-1 两套128倍插值架构及最终实板端点"])
    paragraph_before(chapter8, "图9-1 全FIR与FIR—CIC架构的演进关系及阶段性验证端点", "Caption")
    paragraph_before(chapter8, "图9-1用于呈现算法架构转折，而非建立跨工具版本的资源百分比。Phase 6保留为全FIR回退点，FIR—CIC路线进一步形成2018.3与2025.2两组相互独立的优化序列。", "Normal")
    paragraph_before(chapter8, "9.3.2 Vivado 2025.2同工具链演进", "Heading 3")
    delete_table_rows(t54, [0, 2, 5, 6, 7, 8, 9, 10, 11])
    for row in t54.rows:
        for cell in row.cells:
            text = cell.text.replace("实板通过", "物理板功能验证结果满足判据").replace("板测", "物理板功能验证")
            text = text.replace("routed状态不变量复用", "布局布线后状态不变量复用")
            text = text.replace("发布", "归档")
            text = text.replace("前一实板回退", "前一物理板验证回退点")
            text = text.replace("24/20/20板级验证端点", "24/20/20物理板验证端点")
            set_cell_text(cell, text, 6.7, row is t54.rows[0])
    paragraph_before(chapter8, "表9-5 Vivado 2025.2同工具链的代表性优化节点", "Caption")
    move_before(chapter8, t54._tbl)
    paragraph_before(chapter8, "Vivado 2025.2下，结构优化使完整系统由292 LUT降至221 LUT，下降71 LUT（24.32%）；在此基础上，Stage2有效位宽由22 bit缩减至20 bit后形成218-LUT端点。前一阶段保持24/22/20位数值配置不变，后一阶段属于有限字长Pareto，因此必须结合SQNR、频响和镜像抑制结果解释。", "Normal")
    move_before(chapter8, figure_assets["图9-4 Vivado 2025.2同工具链LUT优化演进"])
    paragraph_before(chapter8, "图9-2 Vivado 2025.2同工具链的LUT优化演进", "Caption")
    paragraph_before(chapter8, "9.3.3 DSP—LUT权衡与优化边界", "Heading 3")
    paragraph_before(chapter8, "固定24/20/20位宽和4个RAMB18E1后，CIC积分器的DSP映射形成三个主要端点：218 LUT/4 DSP、239 LUT/3 DSP与268 LUT/2 DSP。DSP减少并不会自动带来更低总面积，而是将部分宽加法和状态逻辑转移到LUT/CARRY结构；因此三点均应视为Pareto端点。", "Normal")
    move_before(chapter8, figure_assets["图9-6 固定4 RAMB18E1条件下的DSP-LUT Pareto与2-DSP候选复验"])
    paragraph_before(chapter8, "图9-3 固定4个RAMB18E1条件下的DSP—LUT Pareto关系", "Caption")
    paragraph_before(chapter8, "表9-6 固定2 DSP与4个RAMB18E1条件下的候选复验", "Caption")
    move_before(chapter8, t58._tbl)
    set_cell_text(t58.cell(1, 0), "2-DSP目标实现基线")
    set_cell_text(t58.cell(1, 3), "快速与完整回归均17/17项满足判据；保留")
    paragraph_before(chapter8, "2-DSP基线占用268 LUT/417 FF，WNS/WHS为+44.389/+0.078 ns。状态合并和舍入状态表示虽保持逐样本比较误差为0 LSB，却分别增加3和20个LUT；共享Fabric加法器候选还破坏吞吐约束。因此在DSP和BRAM固定时，现有268-LUT实现是已考察候选集合中的资源边界，并不构成全局最优性证明。", "Normal")
    paragraph_before(chapter8, "9.4 同口径资源审计与版本取舍", "Heading 2")
    paragraph_before(chapter8, "9.4.1 218/239版本的核心OOC与完整系统比较", "Heading 3")
    paragraph_before(chapter8, "为获得边界一致的核心资源数据，本文从218-LUT/4-DSP与239-LUT/3-DSP两个不可变版本标签导出RTL快照，在相同Vivado 2025.2、器件、核心顶层、源清单、24/20/20位定点参数、边界时钟及实现指令下执行模块外综合与实现（out-of-context，OOC）；唯一变量为与各版本匹配的CIC积分器DSP模式。每个版本在相同实现设置下重复两次，资源结果一致。", "Normal")
    paragraph_before(chapter8, "表9-7 218-LUT与239-LUT版本的同口径核心OOC与完整系统资源", "Caption")
    move_before(chapter8, t59._tbl)
    set_cell_text(t59.cell(2, 0), "完整板级布局布线后")
    move_before(chapter8, figure_assets["图9-7 218-LUT与239-LUT板测版本的同口径资源比较"])
    paragraph_before(chapter8, "图9-4 218-LUT与239-LUT端点的同口径核心OOC及完整系统资源比较", "Caption")
    paragraph_before(chapter8, "在同口径OOC中，218版与239版分别占用190 LUT/279 FF/4 DSP/3 RAMB18E1和212 LUT/302 FF/3 DSP/3 RAMB18E1；239版相对218版的差分为+22 LUT、+23 FF、−1 DSP和0 RAMB18E1。对应完整系统差分为+21 LUT、+23 FF、−1 DSP和0 RAMB18E1。辅助物理归属支持差异主要集中于CIC映射，但扁平化归属不解释为RTL模块的独立面积，也不与OOC总量相减推导外围资源。", "Normal")
    paragraph_before(chapter8, "两版OOC的内部寄存器路径均满足约束，且DRC Error为0；零I/O延迟边界模型均存在同一条−0.845 ns端口到寄存器保持时间违例，故OOC结果仅用于资源比较，不能表述为整体OOC时序收敛。完整接口时序采用板级布局布线后报告，218版和239版的WNS/WHS分别为+45.279/+0.079 ns与+44.703/+0.079 ns。", "Normal")
    paragraph_before(chapter8, "9.4.2 版本适用范围与实现选择", "Heading 3")
    # Retain the concise parts of the original matrix.
    delete_table_rows(t60, [0, 1, 2, 3, 7])
    for row in t60.rows:
        for cell in row.cells:
            text = cell.text.replace("taps", "个抽头").replace("最低LUT", "LUT占用较低")
            text = text.replace("最终实板", "物理板验证")
            set_cell_text(cell, text, 6.9, row is t60.rows[0])
    paragraph_before(chapter8, "表9-8 三种架构的适用范围与资源取向", "Caption")
    move_before(chapter8, t60._tbl)
    paragraph_before(chapter8, "三种资源端点采用相同24/20/20位数值配置和4个RAMB18E1：218 LUT/4 DSP侧重通用逻辑最小化，239 LUT/3 DSP提供折中，268 LUT/2 DSP满足DSP约束。218版与239版已完成物理板功能验证；268版完成快速与完整RTL回归、布局布线后实现验证及配置位流生成。竞赛报告将218-LUT/4-DSP端点作为LUT优先的主要资源基准，同时保留239-LUT/3-DSP板级验证端点和268-LUT/2-DSP目标实现基线作为可复核的Pareto选择。", "Normal")
    paragraph_before(chapter8, "由此，版本选择应由系统资源约束决定：LUT受限时选择218-LUT端点，DSP与LUT需要平衡时选择239-LUT端点，DSP受限时选择268-LUT基线。上述结论限定于XC7A35T、本文时钟与接口约束以及已覆盖候选集合，不将单一资源最小值表述为无条件最优。", "Normal")

    # Remove the original chapters after selected XML assets have been moved.
    remove_block(chapter8._p, chapter10._p)

    # Update two theoretical citation anchors to use the literature review more coherently.
    polyphase = find_contains(doc, "把128分解为若干小倍率")
    set_paragraph_text(polyphase, polyphase.text.replace("[2-3]", "[11,13-14]") if "[2-3]" in polyphase.text else polyphase.text)
    cic = find_contains(doc, "CIC插值器由低速梳状器")
    if "[12]" in cic.text:
        set_paragraph_text(cic, cic.text.replace("[12]", "[9,12-14]"))

    normalize_references(doc)
    normalization_counts = academic_normalization(doc)
    targeted_academic_rewrites(doc)

    appendix_heading_rewrites = {
        "C.2 RTL仿真PASS关键字与覆盖结论": "C.2 RTL仿真判据与覆盖结果",
        "C.3.1 Release向量规模与输出点数修正": "C.3.1 完整回归向量规模与输出点数核正",
        "C.5 2-DSP生产配置的发布复验与优化边界": "C.5 2-DSP目标实现配置的独立复验与优化边界",
    }
    appendix_targets = {}
    for old, new in appendix_heading_rewrites.items():
        heading_matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == old]
        if not heading_matches and "PASS关键字" in old:
            heading_matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == old.replace("PASS关键字", "日志终止标志“PASS”")]
        if len(heading_matches) != 1:
            raise RuntimeError(f"Expected one appendix heading {old!r}; found {len(heading_matches)}")
        heading = heading_matches[0]
        appendix_targets[old] = (heading, new)
        set_heading_text_preserve_bookmarks(heading, new)
        toc_matches = [p for p in doc.paragraphs if p.style.name.lower().startswith("toc") and p.text.startswith(old)]
        if not toc_matches and "PASS关键字" in old:
            toc_matches = [p for p in doc.paragraphs if p.style.name.lower().startswith("toc") and p.text.startswith(old.replace("PASS关键字", "日志终止标志“PASS”"))]
        if len(toc_matches) > 1:
            raise RuntimeError(f"Expected at most one cached TOC entry for {old!r}; found {len(toc_matches)}")
        if toc_matches:
            rewrite_toc_visible_title(toc_matches[0], new)

    # Insert a cached TOC entry for 8.2; all cached page numbers remain marked for refresh.
    # Replace stale chapter 8/9 cached entries with the new top-level outline.
    toc_paragraphs = list(doc.paragraphs)
    toc8 = next(p for p in toc_paragraphs if p.style.name.lower().startswith("toc") and p.text.startswith("第8章"))
    toc9 = next(p for p in toc_paragraphs if p.style.name.lower().startswith("toc") and p.text.startswith("第9章"))
    toc10 = next(p for p in toc_paragraphs if p.style.name.lower().startswith("toc") and p.text.startswith("第10章"))
    # Remove cached entries between chapter 8 and chapter 10, then add clean chapter 8/9 entries.
    start = toc8._p
    element = start
    while element is not None and element is not toc10._p:
        following = element.getnext()
        element.getparent().remove(element)
        element = following
    toc_anchor = toc10
    cached_entries = [
        ("第8章 历史厂商IP基线与手写RTL的口径化比较", "47", "toc 1"),
        ("8.1 比较目的与对象", "47", "toc 2"),
        ("8.2 统一比较口径与证据边界", "47", "toc 2"),
        ("8.3 历史IP与手写RTL的同口径结果", "48", "toc 2"),
        ("8.4 结果讨论与适用边界", "50", "toc 2"),
        ("8.5 本章小结", "50", "toc 2"),
        ("第9章 架构演进、资源权衡与实现选择", "51", "toc 1"),
        ("9.1 评价准则与证据层级", "51", "toc 2"),
        ("9.2 候选架构的横向比较", "52", "toc 2"),
        ("9.3 架构演进与实现优化", "55", "toc 2"),
        ("9.4 同口径资源审计与版本取舍", "59", "toc 2"),
    ]
    toc1_template = toc10
    toc2_template = next(paragraph for paragraph in doc.paragraphs if paragraph.style.name.lower() == "toc 2")
    next_heading_bookmark_id = max(int(node.get(qn("w:id"))) for node in doc.element.body.iter(qn("w:bookmarkStart")) if (node.get(qn("w:id")) or "").isdigit()) + 1
    for offset, (title, page, style) in enumerate(cached_entries):
        template = toc1_template if style == "toc 1" else toc2_template
        bookmark_name = f"_TocAcademicRevision89{offset:02d}"
        heading = find_exact(doc, title)
        add_toc_bookmark(heading, next_heading_bookmark_id + offset, bookmark_name)
        entry = make_toc_entry(template, title, page, bookmark_name)
        toc_anchor._p.addprevious(entry)

    # Academic title rewrites preserve bookmarks; assert every cached PAGEREF target remains valid.
    document_xml = doc.element.xml
    bookmark_names = set(re.findall(r'w:name="(_Toc[^"]+)"', document_xml))
    for target in re.findall(r'PAGEREF (_Toc\S+)', document_xml):
        if target in bookmark_names:
            continue
        legacy_target = next((p for p in doc.paragraphs if p.style.name.startswith("Heading") and "RTL仿真" in p.text and "覆盖" in p.text), None)
        if target == "_Toc237466528" and legacy_target is not None:
            next_id = max(int(node.get(qn("w:id"))) for node in doc.element.body.iter(qn("w:bookmarkStart")) if (node.get(qn("w:id")) or "").isdigit()) + 1
            add_toc_bookmark(legacy_target, next_id, target)
            bookmark_names.add(target)
            continue
        raise RuntimeError(f"Unresolved TOC PAGEREF target: {target}")

    # Make Word/WPS rebuild TOC fields and pagination on open.
    doc.save(output)
    set_update_fields(output)
    with zipfile.ZipFile(output, "r") as archive:
        bad_member = archive.testzip()
        member_count = len(archive.namelist())
    if bad_member is not None:
        raise RuntimeError(f"Corrupt DOCX member: {bad_member}")

    return {
        "source": str(source),
        "output": str(output),
        "backup": str(backup),
        "source_sha256": sha256(source),
        "output_sha256": sha256(output),
        "zip_member_count": member_count,
        "academic_normalization_counts": normalization_counts,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--backup", required=True, type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    args = parser.parse_args()
    summary = build_revised_document(args.input, args.output, args.backup)
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    args.manifest.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
