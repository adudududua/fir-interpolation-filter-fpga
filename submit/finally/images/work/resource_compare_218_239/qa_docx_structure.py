from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_has_page_break(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:br[@w:type='page']"))


def paragraph_has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: qa_docx_structure.py <docx> <output.json>")
    docx_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    doc = Document(docx_path)

    findings: list[dict[str, object]] = []
    blank_run_threshold = 36
    blank_run_start = None
    blank_count = 0
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text and not paragraph_has_drawing(paragraph) and not paragraph_has_page_break(paragraph):
            if blank_run_start is None:
                blank_run_start = index
            blank_count += 1
        else:
            if blank_count >= blank_run_threshold:
                findings.append({"type": "excessive_blank_paragraphs", "start": blank_run_start, "count": blank_count})
            blank_run_start = None
            blank_count = 0
        if text and re.search(r"(?:差不多|大概|随便|显然|完美|绝对最优|盲目|搞|砍|暴力|应该没问题)", text):
            findings.append({"type": "informal_or_overclaim", "paragraph": index, "text": text})
    if blank_count >= blank_run_threshold:
        findings.append({"type": "excessive_blank_paragraphs", "start": blank_run_start, "count": blank_count})

    expected_headings = {
        "9.4.3 218-LUT与239-LUT板测版本的同口径核心OOC与整板资源比较",
        "9.4.4 三种方案适用条件与两级判据",
        "9.4.5 最终配置、发布标签与结论",
    }
    present = {paragraph.text.strip() for paragraph in doc.paragraphs}
    for heading in sorted(expected_headings - present):
        findings.append({"type": "missing_heading", "heading": heading})

    captions = [p.text.strip() for p in doc.paragraphs if p.style and p.style.name == "Caption"]
    for caption in (
        "表9-16 218-LUT与239-LUT板测版本的同口径核心OOC与完整系统资源比较",
        "图9-7 218-LUT与239-LUT板测版本的同口径资源比较",
        "表9-17 三种方案完整优势—劣势—适用条件矩阵",
        "表9-18 Vivado 2025.2最终系统配置与发布证据",
    ):
        if caption not in captions:
            findings.append({"type": "missing_caption", "caption": caption})

    alt_text_matches = 0
    for paragraph in doc.paragraphs:
        for doc_pr in paragraph._p.xpath(".//wp:docPr"):
            descr = doc_pr.get("descr", "")
            if "218-LUT" in descr and "239-LUT" in descr and "资源" in descr:
                alt_text_matches += 1
    if alt_text_matches != 1:
        findings.append({"type": "resource_figure_alt_text_count", "count": alt_text_matches})

    target_table = None
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "统计边界":
            target_table = table
            break
    if target_table is None:
        findings.append({"type": "missing_resource_table"})
    else:
        last_row = [cell.text.strip() for cell in target_table.rows[-1].cells]
        expected_last_row = [
            "完整板级post-route",
            "218 LUT / 365 FF\n4 DSP / 4 RAMB18E1",
            "239 LUT / 388 FF\n3 DSP / 4 RAMB18E1",
            "+21 LUT / +23 FF\n−1 DSP / 0 RAMB18E1",
        ]
        if last_row != expected_last_row:
            findings.append({"type": "resource_table_total_mismatch", "actual": last_row})

    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml")
        # The template uses w:instrText field codes for the TOC/captions;
        # detect actual tracked-change elements only, not the w:ins prefix.
        if re.search(rb"<w:ins(?:\s|>)", xml) or re.search(rb"<w:del(?:\s|>)", xml):
            findings.append({"type": "tracked_changes_present"})

    report = {
        "file": str(docx_path),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "alt_text_matches": alt_text_matches,
        "findings": findings,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if findings:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
