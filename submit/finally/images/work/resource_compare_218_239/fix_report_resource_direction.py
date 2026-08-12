from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document


OLD_TEXT = (
    "整板差分为−21 LUT、−23 FF和+1 DSP，而CIC局部差分为−23 LUT、−23 FF和+1 DSP；"
    "其余−1/+1/+2 LUT来自FIR1级、2→4倍桥接量化及核心—外围接口共享逻辑的物理映射重组，"
    "合计抵消2 LUT。两版FIR1、FIR2/3、系数/历史BRAM、测试ROM、按键与模式控制及板级控制的寄存器"
    "数量均不变。时序方面，218版和239版的WNS分别为+45.279 ns和+44.703 ns，WHS均为+0.079 ns，"
    "DRC Error均为0；因此，218版的资源下降未以时序或存储资源退化为代价。需要指出，跨层次共享LUT的"
    "数目依赖当前综合与布局布线结果，表中数据用于复现这两个已签核检查点，不外推为独立RTL模块的通用面积常数。"
)

NEW_TEXT = (
    "整板差分为−21 LUT、−23 FF和+1 DSP，而CIC局部差分为−23 LUT、−23 FF和+1 DSP。"
    "由239版切换至218版时，FIR1级、2→4倍桥接量化及核心—外围接口共享逻辑的物理LUT分别变化"
    "+1、−1和+2，合计产生+2 LUT偏移，抵消了CIC局部减少23 LUT中的2 LUT。两版FIR1、FIR2/3、系数/历史BRAM、"
    "测试ROM、按键与模式控制及板级控制的寄存器数量均不变。时序方面，218版和239版的WNS分别为"
    "+45.279 ns和+44.703 ns，WHS均为+0.079 ns，DRC Error均为0；因此，218版的资源下降未以时序或存储资源退化"
    "为代价。需要指出，跨层次共享LUT的数目依赖当前综合与布局布线结果，表中数据用于复现这两个"
    "已签核检查点，不外推为独立RTL模块的通用面积常数。"
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    parser.add_argument("--backup", required=True, type=Path)
    args = parser.parse_args()
    doc = Document(args.document)
    matches = [p for p in doc.paragraphs if p.text == OLD_TEXT]
    if len(matches) != 1:
        raise SystemExit(f"expected exactly one target paragraph, found {len(matches)}")
    args.backup.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.document, args.backup)
    paragraph = matches[0]
    style = paragraph.style
    alignment = paragraph.alignment
    paragraph.clear()
    paragraph.add_run(NEW_TEXT)
    paragraph.style = style
    paragraph.alignment = alignment
    doc.save(args.document)


if __name__ == "__main__":
    main()
