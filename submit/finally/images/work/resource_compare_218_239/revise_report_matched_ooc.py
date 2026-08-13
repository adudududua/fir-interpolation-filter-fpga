from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


def copy_font(src, dst) -> None:
    if src is None:
        return
    for attr in ("name", "size", "bold", "italic", "underline"):
        value = getattr(src.font, attr, None)
        if value is not None:
            setattr(dst.font, attr, value)
    if src.font.name:
        dst._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), src.font.name)


def set_paragraph(paragraph, text: str) -> None:
    source_run = paragraph.runs[0] if paragraph.runs else None
    style = paragraph.style
    alignment = paragraph.alignment
    paragraph.clear()
    run = paragraph.add_run(text)
    copy_font(source_run, run)
    paragraph.style = style
    paragraph.alignment = alignment


def find_contains(doc: Document, fragment: str):
    matches = [p for p in doc.paragraphs if fragment in p.text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph containing {fragment!r}; found {len(matches)}")
    return matches[0]


def set_cell(cell, text: str, font_size: float = 7.4, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(font_size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def replace_resource_table(doc: Document) -> None:
    old = next((t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip() == "模块/统计类别"), None)
    if old is None:
        raise RuntimeError("Old resource-attribution table not found")
    rows = [
        ["统计边界", "218-LUT/4-DSP版", "239-LUT/3-DSP版", "差分（239−218）"],
        ["独立滤波器核心OOC", "190 LUT / 279 FF\n4 DSP / 3 RAMB18E1", "212 LUT / 302 FF\n3 DSP / 3 RAMB18E1", "+22 LUT / +23 FF\n−1 DSP / 0 RAMB18E1"],
        ["完整板级post-route", "218 LUT / 365 FF\n4 DSP / 4 RAMB18E1", "239 LUT / 388 FF\n3 DSP / 4 RAMB18E1", "+21 LUT / +23 FF\n−1 DSP / 0 RAMB18E1"],
    ]
    new = doc.add_table(rows=len(rows), cols=4)
    new.style = old.style
    new.autofit = False
    widths = [1.28, 1.68, 1.68, 1.58]
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            set_cell(new.cell(ri, ci), value, 7.1 if ri else 7.3, ri == 0)
            new.cell(ri, ci).width = Inches(widths[ci])
    header_properties = new.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    new._tbl.getparent().remove(new._tbl)
    old._tbl.addprevious(new._tbl)
    old._tbl.getparent().remove(old._tbl)


def replace_resource_figure(doc: Document, figure: Path) -> None:
    caption = find_contains(doc, "图9-7 218-LUT与239-LUT板测版本")
    element = caption._p.getprevious()
    picture = None
    while element is not None:
        if element.tag == qn("w:p") and element.xpath(".//w:drawing"):
            picture = Paragraph(element, caption._parent)
            break
        element = element.getprevious()
    if picture is None:
        raise RuntimeError("Resource comparison figure paragraph not found")
    picture.clear()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(0)
    picture.paragraph_format.space_after = Pt(0)
    run = picture.add_run()
    run.add_picture(str(figure), width=Inches(6.15))
    alt = "218-LUT与239-LUT板测版本的同口径核心OOC和完整板级资源比较，并注明统计边界与OOC边界时序限定"
    for doc_pr in picture._p.xpath(".//wp:docPr"):
        doc_pr.set("name", "218/239同口径资源比较")
        doc_pr.set("title", "218-LUT与239-LUT同口径资源比较")
        doc_pr.set("descr", alt)


def update_appendix_table(doc: Document) -> None:
    target = next((t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip() == "步骤" and any("核心OOC" in r.cells[0].text for r in t.rows)), None)
    if target is None:
        raise RuntimeError("Appendix reproduction table not found")
    row = next(r for r in target.rows if "核心OOC" in r.cells[0].text)
    set_cell(row.cells[1], "submit/finally/images/work/resource_compare_218_239/run_matched_core_ooc_2025_2.tcl", 7.0)
    set_cell(row.cells[2], "mode2：190 LUT/279 FF/4 DSP；mode1：212 LUT/302 FF/3 DSP；均为3 RAMB18E1，相同设置重复2/2一致", 7.0)


def set_update_fields_on_open(path: Path) -> None:
    temp = path.with_suffix(".updatefields.tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/settings.xml":
                marker = b"<w:updateFields"
                if marker in data:
                    import re

                    data = re.sub(rb'<w:updateFields[^>]*/>', b'<w:updateFields w:val="true"/>', data, count=1)
                else:
                    data = data.replace(b"</w:settings>", b'<w:updateFields w:val="true"/></w:settings>')
            target.writestr(item, data)
    temp.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--figure", required=True, type=Path)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()
    if args.backup:
        args.backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(args.input, args.backup)
    doc = Document(args.input)
    replacements = {
        "最终实现证据来自Vivado 2025.2": "最终实现证据来自Vivado 2025.2对XC7A35T-FGG484-2完整板级工程的综合、布局布线和bitstream生成。正式归档目录为tools/vivado_2025_2/results/20260811_231722；资源、时序、功耗、CDC、DRC和路由状态分别引用对应的实现后报告。完整系统资源取自板级post-route报告；插值滤波器核心资源取自不可变板测标签快照的同流程OOC post-route，二者作为相互独立的统计边界，不相减推导外围资源。",
        "正式完整板级post-route占用218 LUT": "正式完整板级post-route占用218 LUT、365 FF、4个DSP48E1、4个RAMB18E1（2 BRAM Tile）、17个I/O和2个MMCM，LUTRAM为0。LUT与FF利用率分别为1.05%和0.88%，DSP与BRAM Tile利用率分别为4.44%和4.00%。与该板测标签的mode2参数精确匹配的独立插值核心OOC为190 LUT、279 FF、4个DSP48E1和3个RAMB18E1（1.5 BRAM Tile），两次空目录实现结果一致；该数值仅用于核心统计。",
        "9.4.3 218-LUT与239-LUT板测版本的模块级资源分布": "9.4.3 218-LUT与239-LUT板测版本的同口径核心OOC与整板资源比较",
        "为解释4-DSP资源优先端点与3-DSP折中端点之间的硬件代价": "为获得边界一致的核心资源数据，本文从218-LUT/4-DSP和239-LUT/3-DSP两个不可变board-pass标签分别导出RTL快照，在相同Vivado 2025.2、XC7A35T-FGG484-2、顶层interp128_all2x_v7_folded_fir_cic_top_ce、源清单、24/20/20 bit定点参数、6.144 MHz边界时钟以及综合与实现指令下执行OOC post-route；唯一变量为与板测版本匹配的CIC_INTEGRATOR_DSP_MODE=2或1。每个版本均在两个空结果目录中按相同实现设置重复执行，用于验证流程可复现性，不作为跨seed统计。完整系统资源另取自对应板级routed checkpoint，两个统计边界互不混用。",
        "表9-16给出主要功能模块的实现后分布": "表9-16给出同口径正式结果。OOC核心边界包括三级2倍FIR、级间桥接与量化、16倍CIC、核心内部控制以及3个RAMB18E1；不包括测试音频ROM、按键与模式控制、时钟管理、DAC/IO及板级包装。218版和239版按相同实现设置重复两次，分别精确复现190/279/4/3与212/302/3/3，其中末项为RAMB18E1数量。",
        "表9-16 218-LUT与239-LUT板测版本的主要模块资源分布": "表9-16 218-LUT与239-LUT板测版本的同口径核心OOC与完整系统资源比较",
        "239-LUT版本采用CIC_INTEGRATOR_DSP_MODE=1": "以239版减去218版，独立核心的正式差分为+22 LUT、+23 FF、−1 DSP48E1和0 RAMB18E1；完整系统的正式差分为+21 LUT、+23 FF、−1 DSP48E1和0 RAMB18E1。对OOC扁平化DCP进行辅助物理归属时，218/239两版各类别（LUT/FF/DSP/RAMB18E1）分别为：FIR1级37/83/1/1与36/83/1/1，FIR2/3级40/73/1/1与38/73/1/1，共享系数存储0/0/0/1与0/0/0/1，桥接量化41/2/0/0与41/2/0/0，CIC相关逻辑38/121/2/0与62/144/1/0，跨级共享LUT为34与35。该辅助归属可闭合到OOC总量并支持差异集中于CIC映射的结构解释，但其结果依赖扁平化及物理实现，不能表述为各RTL模块的独立面积。",
        "图9-7 218-LUT与239-LUT板测版本的实现后模块资源归属": "图9-7 218-LUT与239-LUT板测版本的同口径资源比较",
        "整板差分为−21 LUT": "两版OOC的内部寄存器路径均满足约束：218版内部WNS/WHS为+150.943/+0.069 ns，239版为+151.611/+0.099 ns，DRC Error均为0。需要限定的是，零I/O延迟的独立边界模型在两版中各产生1条相同的stage3_compensated_mode端口到寄存器hold违例（−0.845 ns）；故上述内部裕量不用于宣称OOC整体时序收敛，完整接口时序采用对应板级post-route签核。218版与239版板级WNS/WHS分别为+45.279/+0.079 ns和+44.703/+0.079 ns，DRC Error均为0。旧的整板物理单元归属仅保留为差异定位证据，不作为模块独立面积，也不与OOC结果相减推导外围资源。",
        "路由后模块审计进一步表明": "同一流程的核心OOC重复实现表明，218-LUT/4-DSP与239-LUT/3-DSP板测端点的核心资源分别为190 LUT/279 FF/4 DSP/3 RAMB18E1和212 LUT/302 FF/3 DSP/3 RAMB18E1。因此，239版相对218版以22个LUT和23个FF换取1个DSP，核心RAMB18E1数量不变；对应完整板级差分为+21 LUT、+23 FF、−1 DSP和0 RAMB18E1。OOC扁平化网表的辅助物理归属及RTL结构共同表明变化主要与16倍CIC的DSP角色映射有关，但该分项不作为模块独立面积。上述同口径总量为DSP—LUT Pareto端点提供了可重复的实现证据。",
    }
    for fragment, text in replacements.items():
        set_paragraph(find_contains(doc, fragment), text)
    replace_resource_table(doc)
    replace_resource_figure(doc, args.figure)
    update_appendix_table(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    set_update_fields_on_open(args.output)


if __name__ == "__main__":
    main()
