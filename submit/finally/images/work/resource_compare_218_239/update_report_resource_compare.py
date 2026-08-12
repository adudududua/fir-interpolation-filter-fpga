from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def copy_font(src, dst) -> None:
    if src is None:
        return
    for attr in ("name", "size", "bold", "italic", "underline"):
        value = getattr(src.font, attr, None)
        if value is not None:
            setattr(dst.font, attr, value)
    if src.font.name:
        dst._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), src.font.name)


def set_paragraph(paragraph, text: str) -> None:
    source_run = paragraph.runs[0] if paragraph.runs else None
    alignment = paragraph.alignment
    style = paragraph.style
    paragraph.clear()
    run = paragraph.add_run(text)
    copy_font(source_run, run)
    paragraph.style = style
    paragraph.alignment = alignment


def insert_paragraph_after(paragraph, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_paragraph._p = new_p
    if style is not None:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def insert_table_after(doc: Document, paragraph, rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = doc.tables[57].style
    table.autofit = False
    table._tbl.getparent().remove(table._tbl)
    paragraph._p.addnext(table._tbl)
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            run = p.add_run(value)
            run.font.name = "宋体"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(7.0 if row_index else 7.3)
            run.bold = row_index == 0
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_properties.append(header_repeat)
    return table


def insert_paragraph_after_table(table, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_paragraph = table._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_paragraph._p = new_p
    if style is not None:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def add_picture_after(paragraph, image_path: Path, width: float, alt: str):
    p = insert_paragraph_after(paragraph, style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    for doc_pr in p._p.xpath(".//wp:docPr"):
        doc_pr.set("name", alt[:80])
        doc_pr.set("title", alt[:120])
        doc_pr.set("descr", alt)
    return p


def add_caption_after(paragraph, text: str):
    p = insert_paragraph_after(paragraph, text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def find_paragraph(doc: Document, exact: str):
    return next(p for p in doc.paragraphs if p.text.strip() == exact)


def insert_resource_section(doc: Document, image_path: Path) -> None:
    anchor = find_paragraph(
        doc,
        "因此，在本轮已覆盖的综合/实现策略、状态表示和受控调度候选范围内，没有获得低于268 LUT且同时保持2 DSP、4 RAMB18E1、原吞吐及位精确行为的完整板级实现。该结论用于界定当前工程的实测优化边界，不构成对全部潜在RTL结构的数学全局最优证明。",
    )
    p = insert_paragraph_after(anchor, "9.4.3 218-LUT与239-LUT板测版本的模块级资源分布", style="Heading 3")
    p = insert_paragraph_after(
        p,
        "为解释4-DSP资源优先端点与3-DSP折中端点之间的硬件代价，本文分别检出两个板测标签，并对其Vivado 2025.2路由后检查点进行只读审计。由于正式工程采用全层次扁平化，常规层次利用率报告仅保留顶层总计；因此，本文以物理LUT BEL、Slice Register、DSP48E1和RAMB18E1为统计单元，先依据实现后单元的RTL源位置归属，再利用寄存器、DSP或BRAM起止点追踪无源位置的扁平化LUT。对同时跨越多个模块的物理LUT单独列为共享逻辑，不作任意拆分。该口径的各项和分别严格回归218/365/4/4与239/388/3/4，其中FF仅指Slice Register，不含封装于OLOGIC中的8个DAC输出寄存器。",
    )
    p = insert_paragraph_after(
        p,
        "表9-16给出主要功能模块的实现后分布。FIR1级在两版中均占用83 FF、1 DSP和1 RAMB18E1，FIR2/3级均占用73 FF、1 DSP和1 RAMB18E1，共享系数存储均占用1个RAMB18E1；测试音频ROM占用余下1个RAMB18E1。两版的FIR与存储结构由此保持稳定，资源差异主要由16倍CIC的DSP映射方式引起。",
    )
    caption = insert_paragraph_after(p, "表9-16 218-LUT与239-LUT板测版本的主要模块资源分布", style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rows = [
        ["模块/统计类别", "218版\nLUT/FF", "239版\nLUT/FF", "DSP\n218/239", "RAMB18\n218/239", "Δ(239−218)"],
        ["FIR1级", "11/83", "10/83", "1/1", "1/1", "−1 LUT"],
        ["FIR2/3级", "20/73", "20/73", "1/1", "1/1", "0"],
        ["FIR共享系数存储", "0/0", "0/0", "0/0", "1/1", "0"],
        ["2→4倍桥接量化", "51/2", "52/2", "0/0", "0/0", "+1 LUT"],
        ["16倍CIC插值", "34/109", "57/132", "2/1", "0/0", "+23 LUT/+23 FF/−1 DSP"],
        ["核心跨级/共享逻辑", "7/0", "7/0", "0/0", "0/0", "0"],
        ["核心—外围接口共享", "59/0", "57/0", "0/0", "0/0", "−2 LUT"],
        ["测试音频ROM", "0/15", "0/15", "0/0", "1/1", "0"],
        ["外围及外围共享逻辑", "36/83", "36/83", "0/0", "0/0", "0"],
        ["完整板级总计", "218/365", "239/388", "4/3", "4/4", "+21 LUT/+23 FF/−1 DSP"],
    ]
    table = insert_table_after(doc, caption, rows, [1.33, 0.88, 0.88, 0.76, 0.80, 1.55])
    p = insert_paragraph_after_table(
        table,
        "239-LUT版本采用CIC_INTEGRATOR_DSP_MODE=1，末级积分保留1个DSP；218-LUT版本采用模式2，新增的DSP48E1执行串行comb角色交换并保存宽位状态。由239版切换到218版后，CIC由57 LUT/132 FF/1 DSP变为34 LUT/109 FF/2 DSP，即以1个DSP换回23个LUT和23个Slice Register。RTL源位置与检查点单元名分别显示新增单元为u_cic_comb_dsp48e1，消失的23个寄存器对应Fabric中的comb_operand_fabric状态。因此，LUT与FF同时下降并非一般性的实现波动，而与CIC宽位减法/状态由Fabric迁入DSP寄存器的数据通路相一致。",
    )
    picture = add_picture_after(
        p,
        image_path,
        6.15,
        "218-LUT和239-LUT板测版本的实现后模块级LUT、FF、DSP与RAMB18资源分布，以及CIC差分解释",
    )
    caption = add_caption_after(picture, "图9-7 218-LUT与239-LUT板测版本的实现后模块资源归属")
    p = insert_paragraph_after(
        caption,
        "整板差分为−21 LUT、−23 FF和+1 DSP，而CIC局部差分为−23 LUT、−23 FF和+1 DSP。由239版切换至218版时，FIR1级、2→4倍桥接量化及核心—外围接口共享逻辑的物理LUT分别变化+1、−1和+2，合计产生+2 LUT偏移，抵消了CIC局部减少23 LUT中的2 LUT。两版FIR1、FIR2/3、系数/历史BRAM、测试ROM、按键与模式控制及板级控制的寄存器数量均不变。时序方面，218版和239版的WNS分别为+45.279 ns和+44.703 ns，WHS均为+0.079 ns，DRC Error均为0；因此，218版的资源下降未以时序或存储资源退化为代价。需要指出，跨层次共享LUT的数目依赖当前综合与布局布线结果，表中数据用于复现这两个已签核检查点，不外推为独立RTL模块的通用面积常数。",
    )

    set_paragraph(find_paragraph(doc, "9.4.3 三种方案适用条件与两级判据"), "9.4.4 三种方案适用条件与两级判据")
    set_paragraph(find_paragraph(doc, "9.4.4 最终配置、发布标签与结论"), "9.4.5 最终配置、发布标签与结论")
    set_paragraph(find_paragraph(doc, "表9-16 三种方案完整优势—劣势—适用条件矩阵"), "表9-17 三种方案完整优势—劣势—适用条件矩阵")
    set_paragraph(find_paragraph(doc, "表9-17 Vivado 2025.2最终系统配置与发布证据"), "表9-18 Vivado 2025.2最终系统配置与发布证据")


def update_conclusion(doc: Document) -> None:
    anchor = find_paragraph(
        doc,
        "在专用资源权衡方面，同一24/20/20数值配置形成218 LUT/4 DSP、239 LUT/3 DSP和268 LUT/2 DSP三个工作点，三者均保持4个RAMB18E1。218-LUT和239-LUT版本已完成物理板功能验证；当前268-LUT生产配置完成Smoke/Release 17/17、完整实现、0 DRC Error和bitstream。固定2-DSP条件下的策略扫描与三项状态/调度候选均未降低LUT，因而保留268-LUT基线，并将该结论限定于本轮已覆盖的试验空间。",
    )
    insert_paragraph_after(
        anchor,
        "路由后模块审计进一步表明，218-LUT与239-LUT板测端点之间的差异集中在16倍CIC：CIC资源由57 LUT/132 FF/1 DSP变为34 LUT/109 FF/2 DSP，即增加1个DSP后减少23个LUT和23个Slice Register；FIR1、FIR2/3及3个核心RAMB18E1保持不变。考虑跨模块物理LUT重组后，完整板级净差为21 LUT和23 FF。该结果为三个DSP—LUT Pareto端点提供了结构层面的解释，而不仅是整板总量比较。",
    )


def update_toc_style(doc: Document) -> None:
    for style_name in ("toc 1", "toc 2", "toc 3"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.size = Pt(9)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.line_spacing = 1


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--figure", required=True, type=Path)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.backup:
        args.backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(args.input, args.backup)
    doc = Document(args.input)
    insert_resource_section(doc, args.figure)
    update_conclusion(doc)
    update_toc_style(doc)
    doc.save(args.output)


if __name__ == "__main__":
    main()
